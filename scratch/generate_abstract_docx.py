import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_abstract_docx(output_path):
    doc = Document()
    
    # 1. Page Setup (Standard 1 inch margins)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        h_p = header.paragraphs[0]
        h_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        h_run = h_p.add_run("MentorHub — Project Abstract & Technical Documentation")
        h_run.font.name = "Calibri"
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = RGBColor(148, 163, 184)

        # Footer
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run("MentorHub Platform | Full-Stack Web Application (FastAPI + React + MongoDB)")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(148, 163, 184)

    # Color Palette
    PRIMARY_NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    SECONDARY_BLUE = RGBColor(37, 99, 235)     # #2563EB
    TEXT_DARK = RGBColor(30, 41, 59)          # #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)      # #64748B

    # Document Header Tag
    p_tag = doc.add_paragraph()
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(4)
    run_tag = p_tag.add_run("PROJECT ABSTRACT & SYSTEM SPECIFICATION")
    run_tag.font.name = "Calibri"
    run_tag.font.size = Pt(9.5)
    run_tag.font.bold = True
    run_tag.font.color.rgb = SECONDARY_BLUE

    # Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("MentorHub: A Smart Mentorship and Community-Driven Collaborative Learning Platform")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(21)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_NAVY

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("An Asynchronous, Full-Stack Web Architecture Integrating 1-on-1 Guidance, Cohort Micro-Communities, Live Virtual Sessions, and Granular Role-Based Access Control")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = TEXT_MUTED

    # Meta Info Card Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_entries = [
        [("Project Domain:", " EdTech / Professional Development Ecosystem"), ("Architectural Pattern:", " Decoupled REST & Real-Time Client-Server")],
        [("Technology Stack:", " React 19, FastAPI (Python), Motor, MongoDB"), ("User Roles:", " Mentees, Vetted Mentors, System Administrators")]
    ]

    col_widths = [Inches(3.25), Inches(3.25)]
    for r_idx, row in enumerate(meta_table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            lbl, val = meta_entries[r_idx][c_idx]
            r1 = p.add_run(lbl)
            r1.font.bold = True
            r1.font.size = Pt(9)
            r1.font.color.rgb = PRIMARY_NAVY
            r2 = p.add_run(val)
            r2.font.size = Pt(9)
            r2.font.color.rgb = TEXT_DARK

    set_table_borders(meta_table, color="E2E8F0", sz="4")

    # Spacing
    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    # ABSTRACT CALLOUT BOX
    abs_box = doc.add_table(rows=1, cols=1)
    abs_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_abs = abs_box.cell(0, 0)
    c_abs.width = Inches(6.5)
    set_cell_background(c_abs, "F1F5F9")
    set_cell_margins(c_abs, top=140, bottom=140, left=160, right=160)
    
    # Left accent border
    tcPr = c_abs._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>\n'
        f'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p_abs_h = c_abs.paragraphs[0]
    p_abs_h.paragraph_format.space_after = Pt(6)
    r_abs_h = p_abs_h.add_run("ABSTRACT")
    r_abs_h.font.name = "Calibri"
    r_abs_h.font.size = Pt(12)
    r_abs_h.font.bold = True
    r_abs_h.font.color.rgb = PRIMARY_NAVY

    p_abs_b = c_abs.add_paragraph()
    p_abs_b.paragraph_format.space_after = Pt(8)
    p_abs_b.paragraph_format.line_spacing = 1.15
    abstract_text = (
        "In the contemporary technological and academic landscape, students, aspiring software engineers, and early-career "
        "professionals routinely encounter substantial challenges in accessing structured career guidance, verified industry mentorship, "
        "and collaborative learning environments. Existing professional networks predominantly emphasize broad, passive networking, "
        "leaving learners stranded across disconnected messaging tools, disparate calendaring links, unverified consultants, and expensive "
        "one-off consultation fees. To address these systemic inefficiencies, this project introduces MentorHub, an end-to-end, full-stack "
        "web-based collaborative mentorship ecosystem engineered to facilitate seamless knowledge transfer and community-driven learning.\n\n"
        "MentorHub integrates a robust three-tier Role-Based Access Control (RBAC) framework supporting Mentees, vetted Mentors, and "
        "Platform Administrators. Through a reactive and modern user interface, mentees can discover domain experts across specialized disciplines "
        "(such as Software Engineering, Cloud Computing, Artificial Intelligence, and Product Design), book personalized 1-on-1 mentorship sessions, "
        "enroll in mentor-curated cohort micro-communities, attend live virtual workshops via integrated Google Meet links, and engage in direct "
        "and community-wide real-time messaging. Mentors are provided with a dedicated control suite to manage community memberships, schedule recurring "
        "or one-off sessions with capacity controls, track subscriber earnings, and evaluate mentee progress. To guarantee high mentorship standards, "
        "the administrative control plane enforces a rigorous vetting mechanism, reviewing prospective mentor credentials, portfolios, and LinkedIn profiles, "
        "while supervising platform analytics, financial streams, and community disputes in real time.\n\n"
        "The system is developed on a cutting-edge, non-blocking asynchronous stack utilizing Python FastAPI, Pydantic v2 schema validation, "
        "and Motor (async MongoDB driver) on the backend, complemented by React 19, Vite, Tailwind CSS v4, and Material-UI (MUI) on the frontend. "
        "Platform integrity and security are guaranteed through JSON Web Token (JWT) stateless authentication, BCrypt password hashing, and "
        "email OTP verification via aiosmtplib. The resulting system establishes an intuitive, scalable, and frictionless hub that bridges the "
        "mentorship gap, lowers access barriers for learners, and fosters sustainable, peer-to-peer technical growth."
    )
    r_abs_b = p_abs_b.add_run(abstract_text)
    r_abs_b.font.name = "Calibri"
    r_abs_b.font.size = Pt(10)
    r_abs_b.font.color.rgb = TEXT_DARK

    p_kw = c_abs.add_paragraph()
    p_kw.paragraph_format.space_after = Pt(0)
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.bold = True
    r_kwh.font.size = Pt(9.5)
    r_kwh.font.color.rgb = PRIMARY_NAVY
    r_kwt = p_kw.add_run(
        "Mentorship Platform, EdTech Ecosystem, Full-Stack Development, FastAPI, React 19, MongoDB (Motor), "
        "Role-Based Access Control (RBAC), Micro-Communities, Cohort-Based Learning, Google Meet Integration, JWT Authentication."
    )
    r_kwt.font.italic = True
    r_kwt.font.size = Pt(9.5)
    r_kwt.font.color.rgb = TEXT_DARK

    # Section divider
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # HELPER FUNCTIONS FOR SECTIONS
    def add_sec_head(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = PRIMARY_NAVY
        return p

    def add_para(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rp = p.add_run(bold_prefix)
            rp.font.name = "Calibri"
            rp.font.size = Pt(10)
            rp.font.bold = True
            rp.font.color.rgb = TEXT_DARK
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = TEXT_DARK
        return p

    # 1. Background & Problem Statement
    add_sec_head("1. Background & Problem Statement")
    add_para(
        "In the digital era, access to tutorials, technical documentation, and online courses is higher than ever; however, learners "
        "frequently struggle with career progression due to the absence of personalized feedback, practical industry advice, and structured accountability. "
        "Traditional platforms and existing market solutions present several core limitations:"
    )
    add_para(" Conventional networking services (e.g., LinkedIn) facilitate broad surface-level connections rather than structured, recurring mentor-mentee engagements.", "• Fragmented Communication: ")
    add_para(" Mentors and mentees are forced to manually orchestrate scheduling (Calendly), messaging (WhatsApp/Slack/Discord), and video meetings (Zoom/Meet), creating friction and scheduling drop-offs.", "• High Tooling Friction: ")
    add_para(" Open-market consulting platforms lack rigorous credential vetting, exposing mentees to substandard guidance, misleading credentials, or financial exploitation.", "• Absence of Rigorous Vetting: ")
    add_para(" Existing platforms prioritize expensive 1-on-1 calls, which are financially out of reach for students. They overlook the collaborative power and affordability of cohort-based micro-communities.", "• Prohibitive Costs & Lack of Community: ")

    # 2. Proposed System & Key Objectives
    add_sec_head("2. Proposed Solution & Project Objectives")
    add_para(
        "MentorHub addresses these challenges by delivering an integrated, end-to-end digital mentorship portal designed with the following core objectives:"
    )
    add_para(" Facilitate structured discovery of vetted industry mentors through categorized domain filtering, verified credentials, and user-submitted ratings and reviews.", "• Intelligent Mentor Discovery: ")
    add_para(" Provide mentors the ability to create and manage domain-specific communities with member capacity rules, monthly subscription fees, and dedicated discussion channels.", "• Micro-Community Infrastructure: ")
    add_para(" Seamlessly schedule, pay for, and attend 1-on-1 sessions and group webinars with integrated Google Meet conference links and automated RSVP tracking.", "• Virtual Session Management: ")
    add_para(" Implement direct 1-to-1 messaging between mentees and mentors alongside public cohort discussion feeds for collaborative problem solving.", "• Real-Time Communication: ")
    add_para(" Ensure platform trust and safety via an administrative oversight panel with credential verification (approving/rejecting mentor applicants) and real-time analytics.", "• Comprehensive Administrative Governance: ")

    # 3. System Architecture & Technology Stack
    add_sec_head("3. System Architecture & Technology Stack")
    add_para(
        "The MentorHub architecture leverages a non-blocking, asynchronous micro-service design pattern that cleanly separates presentation, "
        "application logic, data persistence, and external service integrations:"
    )

    tech_table = doc.add_table(rows=5, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    t_headers = ["Tier / Layer", "Technologies Employed", "Architectural Highlights & Capabilities"]
    t_widths = [Inches(1.4), Inches(2.0), Inches(3.1)]

    for c_idx, cell in enumerate(tech_table.rows[0].cells):
        cell.width = t_widths[c_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t_headers[c_idx])
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    tech_rows = [
        ("Frontend Presentation", "React 19, Vite, Tailwind CSS v4, Material-UI (MUI), Lucide React", "Ultra-responsive single page application (SPA), dynamic role-based dashboards, custom toast notifications, and modular UI components."),
        ("Backend Services", "Python FastAPI, Uvicorn ASGI, Pydantic v2", "Asynchronous non-blocking REST APIs, strict request/response data validation, automatic OpenAPI / Swagger documentation, and high-concurrency request handling."),
        ("Database Layer", "MongoDB with Motor (Async Driver), PyMongo", "NoSQL document store optimized for flexible mentor profiles, nested community structures, session RSVPs, and real-time direct chat transcripts."),
        ("Authentication & Security", "JWT (Python-Jose/PyJWT), BCrypt, aiosmtplib (Email OTP)", "Stateless token-based authorization, granular role guardrails (Mentee, Mentor, Admin), secure password hashing, and two-factor OTP email verification.")
    ]

    for r_idx, row_data in enumerate(tech_rows):
        row = tech_table.rows[r_idx + 1]
        bg_col = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            cell.width = t_widths[c_idx]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(row_data[c_idx])
            r.font.size = Pt(9)
            r.font.color.rgb = TEXT_DARK
            if c_idx == 0:
                r.font.bold = True

    set_table_borders(tech_table, color="CBD5E1", sz="4")
    doc.add_paragraph().paragraph_format.space_before = Pt(8)

    # 4. Major Functional Modules
    add_sec_head("4. Major Functional Modules")
    add_para(" User registration with email OTP verification, mentor directory with domain filters, profile editing, 1-on-1 booking requests, micro-community enrollment, session RSVP, direct messaging, and review submission.", "• Mentee Module: ")
    add_para(" Multi-step mentor application, customizable bio and achievements, cohort community builder, live session scheduler (with Google Meet links and pricing), mentee subscriber management, and earnings tracking.", "• Mentor Module: ")
    add_para(" Centralized administrative cockpit featuring real-time platform metrics (total revenue, active users, session counts), mentor application vetting (approve/reject), community oversight, and review moderation.", "• Admin Operations: ")
    add_para(" Creation of thematic learning micro-communities with subscription fees, community rules, and dedicated group discussion boards for peer learning.", "• Community & Cohorts: ")
    add_para(" Asynchronous direct 1-to-1 messaging between mentees and mentors with unread badges, timestamping, and chat history persistence.", "• Real-Time Messaging: ")
    add_para(" Integrated payment workflows supporting community monthly subscriptions, paid session access passes, and transactional ledger tracking.", "• Financial Engine: ")

    # 5. Software & Hardware Specifications
    add_sec_head("5. System Environment & Specifications")
    add_para(" Modern Web Browser (Google Chrome 110+, Mozilla Firefox, Microsoft Edge), Stable Internet Connection (≥ 5 Mbps for video sessions), 4GB RAM minimum.", "• Client Requirements: ")
    add_para(" Python 3.10+, Node.js 18+, MongoDB Server 6.0+, Vite Build Tool, FastAPI Web Server, Git Version Control.", "• Server & Development Stack: ")

    # 6. Conclusion & Future Roadmap
    add_sec_head("6. Conclusion & Future Enhancements")
    add_para(
        "MentorHub effectively bridges the crucial gap between ambition and experience in professional and technical domains. "
        "By consolidating mentor discovery, verified vetting, community cohorts, live video conferencing, and direct communication under one unified web ecosystem, "
        "the platform eliminates operational friction and fosters high-impact collaborative learning."
    )
    add_para(
        "Future enhancements planned for the platform include AI-driven mentor-mentee matching based on resume embeddings and skill gaps, "
        "native browser-based WebRTC video conferencing with built-in code pair programming and whiteboards, automated certificate issuance upon community completion, "
        "and mobile applications developed with React Native for iOS and Android platforms."
    )

    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    out_path = r"d:\MentorHub\MentorHub\MentorHub_Project_Abstract.docx"
    build_abstract_docx(out_path)
